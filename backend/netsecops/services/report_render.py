"""Rendering a frozen report into a downloadable file (FR-RPT-03).

Rendering never re-reads the estate. It takes the stored `content` and formats it, so
every format of one report says the same thing — and the report's `content_hash`
identifies all of them, because they are one report.

**Tabular output needs a per-template decision, and that is why this is a table rather
than a generic flattener.** A report's content is nested, and there is no single correct
way to flatten it: the executive summary's useful table is one row per device, the
exceptions register's is one row per exception, the compliance report's is one per
control. A generic flattener would produce `top_devices.0.hostname` column headers,
which is a machine-readable rendering of a human's spreadsheet and useful to nobody.

So each template names the collection that becomes rows and the columns that matter.
Columns are named explicitly rather than taken from the first row's keys: a row missing
an optional field would otherwise silently drop that column for every row, and the
reader would never know a field had existed.

**A template with no tabular projection is refused, not returned blank.** An empty
spreadsheet reads as "no findings", which is the most expensive misreading this module
can produce. Such a report is still downloadable as JSON and PDF, both of which can
carry a nested document honestly.

**The caveats travel into every format.** A report's `caveats` block says which
conclusions were *not* drawn — external zones guessed rather than supplied, no KEV feed
ingested, an AAA correlation with no servers examined. In the console those sit beside
the numbers; in a file mailed onward they are the first thing lost, so CSV writes them
into the provenance header, XLSX gets its own sheet and PDF its own section.
"""

from __future__ import annotations

import csv
import io
import json
from typing import Any

from netsecops.core.errors import ValidationProblem
from netsecops.db.models.reporting import Report, ReportFormat, ReportStatus, ReportTemplate

#: template -> (key in content holding the rows, ordered column names)
TABLE_PROJECTIONS: dict[str, tuple[str, tuple[str, ...]]] = {
    ReportTemplate.EXECUTIVE_SUMMARY.value: (
        "top_devices",
        ("hostname", "mgmt_ip", "platform", "criticality", "critical", "high", "findings"),
    ),
    ReportTemplate.EXCEPTIONS_REGISTER.value: (
        "exceptions",
        (
            "check_id",
            "scope",
            "device_id",
            "approver",
            "expires_at",
            "status",
            "expired_at_generation",
            "justification",
        ),
    ),
    ReportTemplate.DEVICE_DETAIL.value: (
        "findings",
        (
            "severity",
            "kind",
            "check_id",
            "cve_id",
            "title",
            "status",
            "first_seen_at",
            "last_seen_at",
            "remediation",
        ),
    ),
    ReportTemplate.GROUP_COMPLIANCE.value: (
        "controls",
        (
            "control",
            "passed",
            "failed",
            # Never merged into `failed`, and never omitted. A control nobody could
            # assess is the most common way a compliance table overstates posture.
            "not_evaluated",
            "not_applicable",
            "errored",
            "never_assessed",
            "checks",
        ),
    ),
    ReportTemplate.FIREWALL_RULEBASE.value: (
        "rules",
        ("order", "name", "action", "enabled", "issues"),
    ),
    ReportTemplate.VULNERABILITY.value: (
        "vulnerabilities",
        (
            "severity",
            "cve_ids",
            "hostname",
            "mgmt_ip",
            "installed_version",
            "fixed_versions",
            "confidence",
            "status",
        ),
    ),
    ReportTemplate.AAA_REVIEW.value: (
        "servers",
        (
            "hostname",
            "product",
            "clients",
            "identity_stores",
            "weak_protocols",
            "admin_mfa_enabled",
            "snapshot_age_days",
            "certificates",
        ),
    ),
    ReportTemplate.DRIFT.value: (
        "devices",
        ("hostname", "mgmt_ip", "state", "severity", "headline", "latest_assessed_at"),
    ),
}

#: The trend report has no single table — it is two sets of totals and the deltas
#: between them, and the honest tabular rendering of that is not a table at all.
NO_TABLE: tuple[str, ...] = (ReportTemplate.TREND.value,)


def content_type_for(fmt: ReportFormat) -> str:
    return {
        ReportFormat.JSON: "application/json",
        ReportFormat.CSV: "text/csv",
        ReportFormat.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ReportFormat.PDF: "application/pdf",
    }[fmt]


def filename_for(report: Report, fmt: ReportFormat) -> str:
    """A filename that says what the artefact is and when it was taken.

    The date is in the name because these files leave the product — they are mailed,
    attached to tickets and filed. `report.csv` in a downloads folder six months later
    is evidence nobody can place.
    """
    stamp = report.generated_at.strftime("%Y%m%d") if report.generated_at else "ungenerated"
    slug = report.template.replace("_", "-")
    return f"netsecops-{slug}-{stamp}-{str(report.id)[:8]}.{fmt.value}"


def render(report: Report, fmt: ReportFormat) -> bytes:
    """Format a stored report. Raises rather than emit a misleading empty file."""
    if report.status != ReportStatus.READY.value:
        raise ValidationProblem(
            f"Report {report.id} is {report.status}. Only a completed report can be "
            "downloaded — a partial one would read as a finished assessment."
        )

    match fmt:
        case ReportFormat.JSON:
            return _json_bytes(report)
        case ReportFormat.CSV:
            return _csv_bytes(report)
        case ReportFormat.XLSX:
            return _xlsx_bytes(report)
        case ReportFormat.PDF:
            return _pdf_bytes(report)

    raise ValidationProblem(f"{fmt} is not a renderable format.")  # pragma: no cover


# ── shared helpers ───────────────────────────────────────────────────────


def _projection(report: Report) -> tuple[str, tuple[str, ...]]:
    projection = TABLE_PROJECTIONS.get(report.template)
    if projection is None:
        raise ValidationProblem(
            f"The {report.template!r} template has no table projection — its content is "
            "not a single table. Download it as JSON or PDF."
        )
    return projection


def _rows_for(report: Report) -> tuple[list[dict[str, Any]], tuple[str, ...]]:
    key, columns = _projection(report)
    rows: list[dict[str, Any]] = list((report.content or {}).get(key) or [])
    return rows, columns


def _cell(value: Any) -> str:
    """Flatten one value for a spreadsheet cell.

    Lists become a semicolon-joined string and dicts their JSON, rather than Python's
    `repr`. A reader opening the file should not have to know what a Python list looks
    like in order to read a list of CVE ids.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, (list, tuple)):
        return "; ".join(_cell(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, sort_keys=True, default=str)
    return str(value)


def _provenance(report: Report) -> list[tuple[str, str]]:
    return [
        ("report", report.title),
        ("generated_at", report.generated_at.isoformat() if report.generated_at else ""),
        ("report_id", str(report.id)),
        ("content_hash", report.content_hash or ""),
    ]


def _caveats(report: Report) -> list[tuple[str, str]]:
    """The `caveats` block as flat label/value pairs, for formats that cannot nest."""
    caveats = (report.content or {}).get("caveats") or {}
    if not isinstance(caveats, dict):
        return []
    return [(str(name), _cell(value)) for name, value in sorted(caveats.items())]


# ── JSON ─────────────────────────────────────────────────────────────────


def _json_bytes(report: Report) -> bytes:
    payload = {
        "report_id": str(report.id),
        "template": report.template,
        "title": report.title,
        "generated_at": report.generated_at.isoformat() if report.generated_at else None,
        # Travels with the file so a recipient can verify the artefact they hold is the
        # one that was generated, without access to the console.
        "content_hash": report.content_hash,
        "parameters": report.parameters,
        "content": report.content,
    }
    return json.dumps(payload, indent=2, sort_keys=True, default=str).encode("utf-8")


# ── CSV ──────────────────────────────────────────────────────────────────


def _csv_bytes(report: Report) -> bytes:
    rows, columns = _rows_for(report)

    buffer = io.StringIO(newline="")
    # A provenance header above the table. Spreadsheet readers tolerate it, and without
    # it a CSV detached from the console is a grid of numbers with no date on it.
    buffer.write(f"# {report.title}\n")
    for label, value in _provenance(report)[1:]:
        buffer.write(f"# {label},{value}\n")
    for label, value in _caveats(report):
        buffer.write(f"# caveat: {label},{value}\n")
    if not rows:
        buffer.write("# no rows — this is an empty result, not a failed one\n")

    writer = csv.DictWriter(buffer, fieldnames=list(columns), extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({column: _cell(row.get(column)) for column in columns})

    return buffer.getvalue().encode("utf-8")


# ── XLSX ─────────────────────────────────────────────────────────────────


def _xlsx_bytes(report: Report) -> bytes:
    """A workbook with the table, its totals and its caveats on separate sheets.

    Separate sheets rather than one: a total stacked above a table is a cell that looks
    like data and sorts like data, and the first person to sort the sheet destroys it.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    rows, columns = _rows_for(report)
    content = report.content or {}

    book = Workbook()
    sheet = book.active
    sheet.title = "Report"
    header_font = Font(bold=True)

    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = header_font
    # Freeze the header so a long finding list stays readable while scrolled.
    sheet.freeze_panes = "A2"

    for row in rows:
        sheet.append([_cell(row.get(column)) for column in columns])

    for index, column in enumerate(columns, start=1):
        widest = max([len(column), *(len(_cell(r.get(column))) for r in rows)])
        sheet.column_dimensions[get_column_letter(index)].width = min(60, max(10, widest + 2))

    if not rows:
        sheet.append(["No rows. This is an empty result, not a failed one."])

    summary = book.create_sheet("Summary")
    summary.append(["Field", "Value"])
    for cell in summary[1]:
        cell.font = header_font
    for label, value in _provenance(report):
        summary.append([label, value])

    totals = content.get("totals")
    if isinstance(totals, dict):
        summary.append([])
        summary.append(["Totals", ""])
        for name, value in totals.items():
            # None stays visibly empty rather than becoming 0: "not assessed" and
            # "none found" must not render as the same cell.
            summary.append([name, "" if value is None else _cell(value)])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 60

    caveats = _caveats(report)
    if caveats:
        sheet_caveats = book.create_sheet("Caveats")
        sheet_caveats.append(["Caveat", "Value"])
        for cell in sheet_caveats[1]:
            cell.font = header_font
        for label, value in caveats:
            sheet_caveats.append([label, value])
        sheet_caveats.column_dimensions["A"].width = 34
        sheet_caveats.column_dimensions["B"].width = 90

    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────

#: Rows beyond this are summarised rather than printed. A 5,000-rule PDF is not a
#: document anybody reads, and the count is stated so nothing looks complete that is not.
PDF_ROW_LIMIT = 200


def _pdf_bytes(report: Report) -> bytes:
    """A filed-and-forwarded document: provenance, totals, caveats, then the table.

    Core fonts only (Helvetica), so no font file ships with the product and no glyph
    licensing question arises. That restricts text to Latin-1, so every string is
    transcoded rather than allowed to raise — a report must not fail to render because
    a device description contains an en dash.
    """
    from fpdf import FPDF

    content = report.content or {}

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 9, _latin1(report.title))
    pdf.ln(1)

    pdf.set_font("Helvetica", "", 9)
    for label, value in _provenance(report)[1:]:
        pdf.cell(0, 5, _latin1(f"{label}: {value}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    totals = content.get("totals")
    if isinstance(totals, dict) and totals:
        _pdf_heading(pdf, "Totals")
        pdf.set_font("Helvetica", "", 10)
        for name, value in totals.items():
            shown = "not assessed" if value is None else _cell(value)
            pdf.cell(0, 5, _latin1(f"{name}: {shown}"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

    caveats = _caveats(report)
    if caveats:
        _pdf_heading(pdf, "Caveats - what this report does not say")
        pdf.set_font("Helvetica", "", 9)
        for label, value in caveats:
            pdf.multi_cell(0, 5, _latin1(f"{label}: {value}"))
        pdf.ln(3)

    if report.template in NO_TABLE:
        _pdf_heading(pdf, "Detail")
        pdf.set_font("Courier", "", 8)
        pdf.multi_cell(0, 4, _latin1(json.dumps(content, indent=2, sort_keys=True, default=str)))
        return bytes(pdf.output())

    rows, columns = _rows_for(report)
    _pdf_heading(pdf, f"Detail ({len(rows)} row{'' if len(rows) == 1 else 's'})")

    if not rows:
        pdf.set_font("Helvetica", "I", 10)
        pdf.cell(0, 6, "No rows. This is an empty result, not a failed one.")
        return bytes(pdf.output())

    usable = pdf.w - pdf.l_margin - pdf.r_margin
    width = usable / len(columns)
    limit = _fit(width)

    pdf.set_font("Helvetica", "B", 8)
    for column in columns:
        pdf.cell(width, 6, _latin1(column)[:limit], border=1)
    pdf.ln()

    pdf.set_font("Helvetica", "", 8)
    for row in rows[:PDF_ROW_LIMIT]:
        for column in columns:
            pdf.cell(width, 5, _latin1(_cell(row.get(column)))[:limit], border=1)
        pdf.ln()

    if len(rows) > PDF_ROW_LIMIT:
        pdf.ln(2)
        pdf.set_font("Helvetica", "I", 9)
        pdf.multi_cell(
            0,
            5,
            _latin1(
                f"{len(rows) - PDF_ROW_LIMIT} further rows are not printed. This PDF is "
                f"a summary of {len(rows)} rows; the CSV, XLSX and JSON downloads carry "
                "all of them."
            ),
        )

    return bytes(pdf.output())


def _pdf_heading(pdf: Any, text: str) -> None:
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 7, _latin1(text), new_x="LMARGIN", new_y="NEXT")


def _fit(width_mm: float) -> int:
    """Roughly how many 8pt characters fit a cell, so text is clipped not overflowed."""
    return max(4, int(width_mm / 1.6))


def _latin1(text: str) -> str:
    """Core PDF fonts are Latin-1. Replace what will not encode rather than raise.

    Losing an en dash from a device description is a cosmetic defect; failing to render
    an auditor's report because of one is not.
    """
    return text.encode("latin-1", "replace").decode("latin-1")


__all__ = [
    "NO_TABLE",
    "PDF_ROW_LIMIT",
    "TABLE_PROJECTIONS",
    "content_type_for",
    "filename_for",
    "render",
]
